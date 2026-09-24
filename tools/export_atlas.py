"""Regenerate the human-readable R&B cheat sheet from the canonical Atlas JSON.

    python tools/export_atlas.py            -> docs/research/RNB_THEORY_ATLAS.xlsx
    python tools/export_atlas.py --csv      -> docs/research/atlas_csv/*.csv

The JSON (auralis/theory/data/rnb_atlas.json) is the source of truth. The
workbook is for reading and review; edits belong in the JSON, then re-export.
Without openpyxl installed, CSVs are written instead.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from auralis.theory import load_atlas  # noqa: E402

OUT_DIR = ROOT / "docs" / "research"


def _j(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_j(v)}" for k, v in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(_j(v) for v in value)
    return str(value)


def sheets(atlas: dict) -> dict[str, tuple[list[str], list[list]]]:
    src_title = {s["id"]: s["title"] for s in atlas["sources"]}
    cite = lambda ids: "; ".join(src_title[i] for i in ids)
    era_name = {e["id"]: e["name"] for e in atlas["eras"]}
    eras_of = lambda ids: ", ".join(era_name.get(i, i) for i in ids)
    return {
        "Era Profiles": (
            ["Era", "BPM tendency", "Harmonic rhythm", "Chord colours", "Groove feel", "Vocal behaviour",
             "Traits", "Status", "Sources"],
            [[e["name"], f"{e['bpm_band'][0]}–{e['bpm_band'][1]}", e["harmonic_rhythm"], _j(e["chord_colors"]),
              _j(e["groove"]), _j(e["vocal"]), "\n".join(e["traits"]), e["status"], cite(e["sources"])]
             for e in atlas["eras"]]),
        "Progression Families": (
            ["ID", "Name", "Roman numerals", "Mode", "Loop", "Cadence / loop type", "Tension",
             "Era affinity", "Section affinity", "Reharmonization options", "Comment", "Status", "Sources"],
            [[p["id"], p["name"], " – ".join(p["roman"]), p["mode"], "yes" if p["loop"] else "no", p["cadence"],
              p["tension"], _j({era_name[k]: v for k, v in p["eras"].items()}), _j(p["sections"]), "\n".join(p["reharm"]),
              p.get("comment", ""), p["status"], cite(p["sources"])]
             for p in atlas["progression_families"]]),
        "Chord Vocabulary": (
            ["Quality", "Function", "Common approach", "Common exit", "Extensions", "Omissions",
             "Voicing / rootless behaviour", "Era affinity", "Status", "Sources"],
            [[c["quality"], c["function"], c["approach"], c["exit"], _j(c["extensions"]), _j(c["omissions"]),
              c["voicing"], eras_of(c["eras"]), c["status"], cite(c["sources"])]
             for c in atlas["chord_vocabulary"]]),
        "Vocal Phrase Patterns": (
            ["ID", "Name", "Section", "Approach", "Contour", "Start degree", "End degree", "Peak degree",
             "Phrase bars", "Pickup", "Syncopation", "Melisma", "Register", "Hook repetition", "Eras", "Status", "Sources"],
            [[v["id"], v["name"], v["section"], _j(v["approach"]), v["contour"], v["start_degree"], v["end_degree"],
              v["peak_degree"], f"{v['phrase_bars'][0]}–{v['phrase_bars'][1]}", v["pickup"], v["syncopation"],
              v["melisma"], v["register"], v["hook_repetition"], eras_of(v["eras"]), v["status"], cite(v["sources"])]
             for v in atlas["vocal_patterns"]]),
        "Groove - Pocket": (
            ["ID", "Name", "Feel", "Eras", "BPM band", "Swing ratio", "Kick / snare / bass / comping offsets (ms)",
             "Quantize strength", "Push / pull", "Pocket width (ms)", "Comment", "Status", "Sources"],
            [[g["id"], g["name"], g["feel"], eras_of(g["eras"]), f"{g['bpm_band'][0]}–{g['bpm_band'][1]}",
              f"{g['swing_ratio'][0]}–{g['swing_ratio'][1]}",
              "; ".join(f"{k} {lo}–{hi}" for k, (lo, hi) in g["offsets_ms"].items()),
              f"{g['quantize_strength'][0]}–{g['quantize_strength'][1]}", g["push_pull"], g["pocket_width_ms"],
              g.get("comment", ""), g["status"], cite(g["sources"])]
             for g in atlas["grooves"]]),
        "Song Evidence": (
            ["ID", "Work", "Artist", "Year", "Dataset", "Abstract observation", "Roman abstraction",
             "Features", "Confidence", "Sources"],
            [[ev["id"], ev["work"], ev.get("artist") or "", ev.get("year") or "", ev.get("dataset") or "",
              ev["observation"], " – ".join(ev["abstraction"].get("roman", [])),
              _j(ev["abstraction"].get("features", [])), ev["confidence"], cite(ev["sources"])]
             for ev in atlas["evidence"]]),
        "Sources": (
            ["ID", "Title", "Author", "Publisher", "Year", "URL", "Comment"],
            [[s["id"], s["title"], s["author"], s["publisher"], s.get("year") or "", s.get("url") or "", s.get("comment", "")]
             for s in atlas["sources"]]),
    }


def write_xlsx(data, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1A1509")
    for name, (header, rows) in data.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="E9CF8A")
            cell.fill = header_fill
        ws.freeze_panes = "B2"
        for col in ws.columns:
            width = min(60, max(10, max(len(str(c.value or "").split("\n")[0]) for c in col) + 2))
            ws.column_dimensions[col[0].column_letter].width = width
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--csv", action="store_true", help="write CSVs instead of a workbook")
    args = parser.parse_args(argv)
    data = sheets(load_atlas())
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    use_csv = args.csv
    if not use_csv:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            print("openpyxl is not installed; writing CSVs instead (pip install openpyxl for the workbook).")
            use_csv = True
    if use_csv:
        folder = OUT_DIR / "atlas_csv"
        folder.mkdir(exist_ok=True)
        for name, (header, rows) in data.items():
            with open(folder / f"{name.replace(' ', '_').lower()}.csv", "w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.writer(fh)
                writer.writerow(header)
                writer.writerows(rows)
        print(f"wrote {len(data)} CSV sheets to {folder}")
    else:
        path = OUT_DIR / "RNB_THEORY_ATLAS.xlsx"
        write_xlsx(data, path)
        print(f"wrote {path} ({', '.join(data)})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
